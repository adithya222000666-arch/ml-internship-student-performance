"""
Script to generate the polished B.Tech Capstone Presentation (.pptx)
Student: Adithya M.
Week 4 Capstone: Student Performance Prediction System
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

DOCS_DIR = Path(__file__).resolve().parent
WEEK4_DIR = DOCS_DIR.parent
OUTPUTS_DIR = WEEK4_DIR / "outputs"
PPTX_PATH = DOCS_DIR / "presentation.pptx"

# Color Palette Constants
COLOR_NAVY = RGBColor(30, 58, 138)       # #1E3A8A - Primary Header
COLOR_BLUE = RGBColor(37, 99, 235)       # #2563EB - Accent Blue
COLOR_SLATE = RGBColor(71, 85, 105)      # #475569 - Muted Text
COLOR_DARK = RGBColor(15, 23, 42)        # #0F172A - Body Text
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_CARD_BG = RGBColor(241, 245, 249)  # #F1F5F9 - Card Background
COLOR_CARD_BORDER = RGBColor(203, 213, 225)
COLOR_EMERALD = RGBColor(5, 150, 105)    # #059669 - Accent Green
COLOR_AMBER = RGBColor(217, 119, 6)      # #D97706 - Warning / Placeholder


def create_blank_slide(prs):
    """Create a clean blank slide."""
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    return slide


def add_slide_header(slide, title_text, category="STUDENT PERFORMANCE PREDICTION SYSTEM"):
    """Add a consistent, professional header across content slides."""
    # Category Tracker
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_BLUE

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_NAVY

    # Thin Divider Line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_CARD_BORDER
    line.line.color.rgb = COLOR_CARD_BORDER


def add_content_card(slide, left, top, width, height, title, items, is_accent=False, is_placeholder=False):
    """Add a structured container card with bullet points."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(254, 243, 199) if is_placeholder else (RGBColor(238, 242, 255) if is_accent else COLOR_CARD_BG)
    shape.line.color.rgb = COLOR_AMBER if is_placeholder else (COLOR_BLUE if is_accent else COLOR_CARD_BORDER)
    shape.line.width = Pt(1.5 if is_accent or is_placeholder else 1)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    tf.margin_bottom = Inches(0.2)

    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.bold = True
    p0.font.size = Pt(14)
    p0.font.color.rgb = COLOR_AMBER if is_placeholder else (COLOR_NAVY if is_accent else COLOR_NAVY)
    p0.space_after = Pt(8)

    for item in items:
        p = tf.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK
        p.space_after = Pt(4)


def generate_presentation():
    prs = Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = create_blank_slide(prs)

    # Background accent bar
    bg_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.2))
    bg_bar.fill.solid()
    bg_bar.fill.fore_color.rgb = COLOR_NAVY
    bg_bar.line.fill.background()

    # Main Title Container
    title_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(2.2))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "MACHINE LEARNING INTERNSHIP CAPSTONE • WEEK 4"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_BLUE
    p_badge.space_after = Pt(12)

    p_main = tf1.add_paragraph()
    p_main.text = "Student Performance Prediction System"
    p_main.font.size = Pt(36)
    p_main.font.bold = True
    p_main.font.color.rgb = COLOR_NAVY
    p_main.space_after = Pt(8)

    p_sub = tf1.add_paragraph()
    p_sub.text = "End-to-End Supervised ML Pipeline, Flask Web Application, Automated QA Suite & Render Cloud Deployment"
    p_sub.font.size = Pt(15)
    p_sub.font.color.rgb = COLOR_SLATE

    # Presenter Information Card
    info_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.5), Inches(11.333), Inches(1.8))
    info_card.fill.solid()
    info_card.fill.fore_color.rgb = COLOR_CARD_BG
    info_card.line.color.rgb = COLOR_BLUE

    tf_info = info_card.text_frame
    tf_info.margin_left = Inches(0.4)
    tf_info.margin_top = Inches(0.25)

    p_st = tf_info.paragraphs[0]
    p_st.text = "Student Presenter: Adithya M.  |  B.Tech Computer Science & Engineering"
    p_st.font.bold = True
    p_st.font.size = Pt(14)
    p_st.font.color.rgb = COLOR_NAVY
    p_st.space_after = Pt(4)

    p_rep = tf_info.add_paragraph()
    p_rep.text = "Repository: github.com/adithya222000666-arch/ml-internship-student-performance"
    p_rep.font.size = Pt(11)
    p_rep.font.color.rgb = COLOR_BLUE
    p_rep.space_after = Pt(4)

    p_date = tf_info.add_paragraph()
    p_date.text = "Academic Evaluation Session: October 2026  |  Final Capstone Presentation"
    p_date.font.size = Pt(11)
    p_date.font.color.rgb = COLOR_SLATE

    # =========================================================================
    # SLIDE 2: PROBLEM STATEMENT
    # =========================================================================
    s2 = create_blank_slide(prs)
    add_slide_header(s2, "Problem Statement: Proactive vs. Reactive Academic Warning")
    add_content_card(s2, 0.8, 1.8, 5.6, 5.0, "The Traditional Academic Challenge", [
        "Conventional student intervention is inherently reactive.",
        "Schools identify struggling students only after midterm marks (G1/G2) or final failures are published.",
        "Late-stage remediation leaves minimal time for effective pedagogical recovery.",
        "High school mathematics failure severely impairs university admission prospects."
    ])
    add_content_card(s2, 6.8, 1.8, 5.7, 5.0, "The Machine Learning Solution", [
        "Eliminate Grade Data Leakage: G1 & G2 are excluded completely to ensure authentic early prediction.",
        "Multi-Factor Profiling: Leverages 30 socio-demographic, familial, and lifestyle indicators.",
        "Advisory Early Warning: Flags students needing mentoring before initial examinations take place.",
        "Decision-Support System: Provides quantifiable estimates (0–20 scale) to guide academic counselors."
    ], is_accent=True)

    # =========================================================================
    # SLIDE 3: PROJECT OBJECTIVES
    # =========================================================================
    s3 = create_blank_slide(prs)
    add_slide_header(s3, "Project Objectives & Engineering Scope")
    add_content_card(s3, 0.8, 1.8, 3.6, 5.0, "1. ML Engineering", [
        "Robust pipeline via Scikit-Learn ColumnTransformer.",
        "Median/mode imputation & standard scaling.",
        "One-hot encoding handling unseen categories.",
        "5-Fold Cross-Validation on training split."
    ])
    add_content_card(s3, 4.8, 1.8, 3.6, 5.0, "2. Full-Stack App & API", [
        "Modern responsive web application in Flask.",
        "All 30 verified non-grade features in UI.",
        "Sample student preset profiles for quick testing.",
        "RESTful API (/health, /predict, /api/schema)."
    ], is_accent=True)
    add_content_card(s3, 8.8, 1.8, 3.7, 5.0, "3. QA & Deployment", [
        "Automated Pytest suite (16 comprehensive tests).",
        "Edge-case, out-of-range & malformed JSON checks.",
        "Non-retraining runtime guarantee.",
        "Production Gunicorn WSGI for Render deployment."
    ])

    # =========================================================================
    # SLIDE 4: DATASET OVERVIEW
    # =========================================================================
    s4 = create_blank_slide(prs)
    add_slide_header(s4, "Dataset Architecture & Feature Distribution")
    add_content_card(s4, 0.8, 1.8, 5.6, 5.0, "UCI Student Performance (student-mat.csv)", [
        "Origin: Secondary schools in Portugal (Gabriel Pereira & Mousinho da Silveira).",
        "Records: 395 student instances, semicolon-delimited CSV.",
        "Target Variable: Final Mathematics Grade (G3), scale 0–20 (passing score >= 10).",
        "Excluded Features: G1 (First Period) & G2 (Second Period) to prevent data leakage.",
        "Retained Features: Exactly 30 verified operational attributes."
    ])
    add_content_card(s4, 6.8, 1.8, 5.7, 5.0, "Feature Groups (30 Input Features)", [
        "Academic & Schooling (8 features): school, studytime, failures, schoolsup, famsup, paid, higher, absences.",
        "Demographics & Living (8 features): sex, age (15-22), address (Urban/Rural), famsize, Pstatus, guardian, traveltime, reason.",
        "Parental Background (4 features): Medu (0-4), Fedu (0-4), Mjob, Fjob.",
        "Social, Lifestyle & Well-being (10 features): internet, romantic, activities, nursery, famrel (1-5), freetime, goout, Dalc, Walc, health."
    ], is_accent=True)

    # =========================================================================
    # SLIDE 5: DATA PREPROCESSING
    # =========================================================================
    s5 = create_blank_slide(prs)
    add_slide_header(s5, "Data Preprocessing: Preventing Data Leakage")
    add_content_card(s5, 0.8, 1.8, 5.6, 5.0, "Numeric Feature Pipeline (13 Features)", [
        "Features: age, Medu, Fedu, traveltime, studytime, failures, famrel, freetime, goout, Dalc, Walc, health, absences.",
        "Step 1: SimpleImputer(strategy='median') - fills missing values using training medians.",
        "Step 2: StandardScaler() - standardizes numeric features (mean = 0, variance = 1).",
        "Leakage Prevention: Preprocessor fitted ONLY on the 80% training partition."
    ])
    add_content_card(s5, 6.8, 1.8, 5.7, 5.0, "Categorical Feature Pipeline (17 Features)", [
        "Features: school, sex, address, famsize, Pstatus, Mjob, Fjob, reason, guardian, schoolsup, famsup, paid, activities, nursery, higher, internet, romantic.",
        "Step 1: SimpleImputer(strategy='most_frequent') - fills missing values with training mode.",
        "Step 2: OneHotEncoder(handle_unknown='ignore', sparse_output=False) - converts nominal categories into binary vectors.",
        "Safety: Unseen categories during live inference are safely zero-filled."
    ], is_accent=True)

    # =========================================================================
    # SLIDE 6: MACHINE LEARNING MODELS
    # =========================================================================
    s6 = create_blank_slide(prs)
    add_slide_header(s6, "Evaluated Supervised Learning Models")
    add_content_card(s6, 0.8, 1.8, 3.6, 5.0, "Linear Regression", [
        "Baseline linear model.",
        "Ordinary Least Squares estimation.",
        "Establishes transparent linear benchmark.",
        "Test MAE: 3.3953",
        "Test RMSE: 4.1957",
        "Test R²: 0.1415"
    ])
    add_content_card(s6, 4.8, 1.8, 3.6, 5.0, "Ridge Regression", [
        "L2 Regularized linear regression (alpha=1.0).",
        "Penalizes large weights to mitigate multi-collinearity.",
        "Improves linear generalization slightly.",
        "Test MAE: 3.3929",
        "Test RMSE: 4.1934",
        "Test R²: 0.1424"
    ])
    add_content_card(s6, 8.8, 1.8, 3.7, 5.0, "Random Forest (Champion)", [
        "Ensemble of 200 decision trees (min_samples_leaf=2).",
        "Captures complex non-linear feature interactions.",
        "Robust against outliers and feature noise.",
        "Test MAE: 3.0364",
        "Test RMSE: 3.8126",
        "Test R²: 0.2911"
    ], is_accent=True)

    # =========================================================================
    # SLIDE 7: EVALUATION METHODOLOGY
    # =========================================================================
    s7 = create_blank_slide(prs)
    add_slide_header(s7, "Model Evaluation & Validation Methodology")
    add_content_card(s7, 0.8, 1.8, 5.6, 5.0, "Reproducible Validation Strategy", [
        "80% Training Set (316 samples) / 20% Test Set (79 samples).",
        "Fixed random_state=42 for deterministic reproduction.",
        "5-Fold Cross-Validation on the training partition.",
        "Cross-validation confirms generalization before final test evaluation."
    ])
    add_content_card(s7, 6.8, 1.8, 5.7, 5.0, "Key Regression Evaluation Metrics", [
        "Mean Absolute Error (MAE): Measures average absolute grade deviation (0–20 scale). Lower is better.",
        "Root Mean Squared Error (RMSE): Penalizes severe deviations quadratically. Lower is better.",
        "Coefficient of Determination (R²): Measures proportion of target variance explained. Higher is better.",
        "Historical Fidelity: Week 4 implementation recomputed and confirmed Week 2 historical baseline metrics exactly."
    ], is_accent=True)

    # =========================================================================
    # SLIDE 8: ACTUAL MODEL RESULTS
    # =========================================================================
    s8 = create_blank_slide(prs)
    add_slide_header(s8, "Actual Experimental Results & Model Selection")

    # Metrics Summary Box
    add_content_card(s8, 0.8, 1.8, 5.2, 5.0, "Model Comparison Benchmarks", [
        "Random Forest selected as Champion Model.",
        "Lowest Test MAE: 3.0364 (vs. 3.3953 for Linear Regression).",
        "Lowest Test RMSE: 3.8126 (vs. 4.1957 for Linear Regression).",
        "Highest Test R²: 0.2911 (vs. 0.1415 for Linear Regression).",
        "5-Fold CV MAE: 2.9343, CV R²: 0.2036.",
        "Top Predictive Drivers: past failures, school absences, student age, going out with friends, parental education."
    ], is_accent=True)

    # Embed Chart 1
    chart1 = OUTPUTS_DIR / "model_comparison_metrics.png"
    if chart1.exists():
        s8.shapes.add_picture(str(chart1), Inches(6.3), Inches(1.8), Inches(6.2), Inches(4.8))

    # =========================================================================
    # SLIDE 9: SYSTEM ARCHITECTURE
    # =========================================================================
    s9 = create_blank_slide(prs)
    add_slide_header(s9, "End-to-End System Architecture")
    add_content_card(s9, 0.8, 1.8, 3.6, 5.0, "1. Offline Training Phase", [
        "Raw Data Ingestion (student-mat.csv).",
        "Feature Selection (drop G1, G2, G3).",
        "Pipeline fitting via ColumnTransformer.",
        "Model serialization to Joblib artifact."
    ])
    add_content_card(s9, 4.8, 1.8, 3.6, 5.0, "2. Model Serving Core", [
        "Singleton pipeline loading at startup.",
        "Zero-latency memory resident model.",
        "model_metadata.json schema validation.",
        "Immutable inference guarantee (no retraining)."
    ], is_accent=True)
    add_content_card(s9, 8.8, 1.8, 3.7, 5.0, "3. Application Presentation", [
        "Flask WSGI web application.",
        "Responsive HTML5, modern CSS3 & JS UI.",
        "RESTful JSON API endpoints.",
        "Production deployment via Gunicorn on Render."
    ])

    # =========================================================================
    # SLIDE 10: WEB APPLICATION INTERFACE
    # =========================================================================
    s10 = create_blank_slide(prs)
    add_slide_header(s10, "Interactive Web Application Interface")
    add_content_card(s10, 0.8, 1.8, 5.6, 5.0, "Front-End UX & Key Features", [
        "Accessible, responsive design tailored for college capstone presentation.",
        "Four logical field groupings covering all 30 input features.",
        "Preset buttons: 'High Performer', 'Average Student', 'At-Risk Student'.",
        "Real-time visual grade gauge (0–20 scale) with performance badges.",
        "Comprehensive advisory disclaimer clarifying statistical estimation."
    ])
    add_content_card(s10, 6.8, 1.8, 5.7, 5.0, "UI Verification & Screenshot Placeholder", [
        "[SCREENSHOT PLACEHOLDER]",
        "Capture live screenshot of the web form with sample student loaded.",
        "Include prediction results card showing the predicted G3 grade, progress bar, and academic standing category.",
        "File path: week4/templates/index.html & static/css/style.css"
    ], is_placeholder=True)

    # =========================================================================
    # SLIDE 11: PREDICTION API
    # =========================================================================
    s11 = create_blank_slide(prs)
    add_slide_header(s11, "RESTful Prediction API Specification")
    add_content_card(s11, 0.8, 1.8, 5.6, 5.0, "API Endpoints Specification", [
        "GET / - Serves interactive web application interface.",
        "GET /health - Returns JSON health status, model readiness, and champion metrics (200 / 503).",
        "GET /api/schema - Exposes all 30 features, data types, and allowed categorical options.",
        "POST /predict - Accepts JSON student attributes; returns predicted grade, category, and disclaimer."
    ])
    add_content_card(s11, 6.8, 1.8, 5.7, 5.0, "Robust Input Validation & Error Handling", [
        "Rejects missing attributes with detailed missing field names (400 Bad Request).",
        "Enforces numeric types and valid numerical ranges (e.g. age: 15–22).",
        "Rejects non-finite values (NaN, Infinity).",
        "Rejects unrecognized categorical inputs.",
        "Never retrains the model during request processing."
    ], is_accent=True)

    # =========================================================================
    # SLIDE 12: MODEL SERIALIZATION
    # =========================================================================
    s12 = create_blank_slide(prs)
    add_slide_header(s12, "Model Serialization & Inference Decoupling")
    add_content_card(s12, 0.8, 1.8, 5.6, 5.0, "Serialized Pipeline Artifact", [
        "Artifact: student_performance_pipeline.joblib (2.45 MB).",
        "Encapsulates both preprocessing (ColumnTransformer) and champion estimator (RandomForestRegressor).",
        "Guarantees that raw incoming JSON data is transformed identically to training data.",
        "Single-file deployment artifact with zero external database dependencies."
    ])
    add_content_card(s12, 6.8, 1.8, 5.7, 5.0, "Schema Metadata Specification", [
        "Metadata: model_metadata.json (9.7 KB).",
        "Preserves exact feature ordering (30 features).",
        "Documents numerical ranges (min, max, median defaults).",
        "Documents allowed categorical string values.",
        "Stores champion benchmark metrics for the /health endpoint."
    ], is_accent=True)

    # =========================================================================
    # SLIDE 13: AUTOMATED TESTING
    # =========================================================================
    s13 = create_blank_slide(prs)
    add_slide_header(s13, "Automated QA & Pytest Verification")
    add_content_card(s13, 0.8, 1.8, 5.6, 5.0, "Pytest Test Suite (week4/tests/test_app.py)", [
        "16 Comprehensive Automated Tests:",
        "• Model artifact loading & health readiness.",
        "• Standard JSON prediction request & schema conformance.",
        "• Form-encoded POST submission support.",
        "• Missing required fields validation (400).",
        "• Invalid numeric types & out-of-range bounds.",
        "• Non-finite numeric rejection (NaN / Inf).",
        "• Invalid category rejection & malformed JSON syntax.",
        "• Output bounding guarantee (0–20 scale).",
        "• Non-retraining model invariance check."
    ])
    add_content_card(s13, 6.8, 1.8, 5.7, 5.0, "Execution Results & Evidence", [
        "Test Status: 16 passed in 2.08 seconds (100% Pass Rate).",
        "[SCREENSHOT PLACEHOLDER]",
        "Capture terminal output of 'pytest week4/tests -v' demonstrating all 16 tests passing.",
        "Confirms enterprise-grade robustness before cloud deployment."
    ], is_accent=True, is_placeholder=False)

    # =========================================================================
    # SLIDE 14: DEPLOYMENT PREPARATION
    # =========================================================================
    s14 = create_blank_slide(prs)
    add_slide_header(s14, "Cloud Deployment Strategy for Render")
    add_content_card(s14, 0.8, 1.8, 5.6, 5.0, "Render Production Configuration", [
        "WSGI Server: Gunicorn multi-worker HTTP server.",
        "Procfile: web: gunicorn app:app --bind 0.0.0.0:$PORT",
        "Runtime: Python 3.11.9 (configured via runtime.txt).",
        "Blueprint: render.yaml with rootDir: week4.",
        "Path Resolution: Relative via Path(__file__).resolve().parent (no Windows path lock-in)."
    ])
    add_content_card(s14, 6.8, 1.8, 5.7, 5.0, "Deployment Verification Checklist", [
        "1. Push repository changes to GitHub main branch.",
        "2. Connect repository to Render Web Service dashboard.",
        "3. Set Root Directory to week4 and branch to main.",
        "4. Set Build Command: pip install -r requirements.txt",
        "5. Set Start Command: gunicorn app:app --bind 0.0.0.0:$PORT",
        "6. Verify live health check: https://<app>.onrender.com/health",
        "[SCREENSHOT PLACEHOLDER: Insert Render Dashboard Screenshot]"
    ], is_accent=True)

    # =========================================================================
    # SLIDE 15: LIMITATIONS & ETHICS
    # =========================================================================
    s15 = create_blank_slide(prs)
    add_slide_header(s15, "Model Limitations & Ethical Considerations")
    add_content_card(s15, 0.8, 1.8, 5.6, 5.0, "Technical & Statistical Limitations", [
        "Non-Grade Signal: Excluding G1 and G2 limits R² to ~0.29. This is an intentional tradeoff to prioritize early warning capability.",
        "Sample Size: 395 records from two regional secondary schools; findings should be generalized cautiously.",
        "Static Survey Data: Does not capture longitudinal attendance fluctuations or mid-semester interventions."
    ])
    add_content_card(s15, 6.8, 1.8, 5.7, 5.0, "Ethical & Educational Principles", [
        "Advisory Instrument Only: Must never be used to deny academic opportunities, streaming, or funding.",
        "Demographic Fairness: Socio-demographic factors reflect cohort patterns and must not pigeonhole individual students.",
        "Human-in-the-Loop: Supports educators and academic counselors—does not replace empathetic mentorship."
    ], is_accent=True)

    # =========================================================================
    # SLIDE 16: CONCLUSION & ROADMAP
    # =========================================================================
    s16 = create_blank_slide(prs)
    add_slide_header(s16, "Conclusion & Future Improvements")
    add_content_card(s16, 0.8, 1.8, 5.6, 5.0, "Capstone Accomplishments", [
        "Complete, production-ready ML system built and verified.",
        "Reproduced Week 2 benchmark results with mathematical precision.",
        "Serialized pipeline artifact with zero data leakage.",
        "Responsive, modern Flask web application.",
        "Robust RESTful JSON prediction API.",
        "16-test automated Pytest suite (100% pass rate).",
        "Production-ready deployment configuration for Render."
    ], is_accent=True)
    add_content_card(s16, 6.8, 1.8, 5.7, 5.0, "Future Engineering Roadmap", [
        "Explainable AI (SHAP / LIME): Provide personalized visual explanations for each student's risk factors.",
        "Prescriptive Recommendations: Suggest actionable study plans based on predicted outcomes.",
        "School Information System (SIS) Integration: Automated batch ingestion of semester attendance data."
    ])

    prs.save(PPTX_PATH)
    print(f"[SUCCESS] Presentation generated at: {PPTX_PATH}")


if __name__ == "__main__":
    generate_presentation()
